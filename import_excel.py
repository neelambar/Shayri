"""Bulk-import poems from an Excel sheet with columns POET, TITLE, TEXT.

TEXT may contain HTML (<p>, <br />, &nbsp; ...); it is converted to plain
lines, with a blank line between paragraphs. A trailing " / <poet>" is
removed from titles. Poems whose text is already in the collection are
skipped, so running the same file twice is safe.

    .venv/bin/python import_excel.py poems.xlsx --dry-run
    SHAYRI_DB=/var/lib/shayri/shayri.db .venv/bin/python import_excel.py poems.xlsx
"""
import argparse
import html
import re
import sys

from openpyxl import load_workbook

import app as shayri
from app import clean, fold

BR = re.compile(r"<br\s*/?>", re.IGNORECASE)
PARA = re.compile(r"</p\s*>\s*<p[^>]*>", re.IGNORECASE)
TAG = re.compile(r"<[^>]+>")


def html_to_text(s):
    s = str(s or "")
    s = PARA.sub("\n\n", s)
    s = BR.sub("\n", s)
    s = TAG.sub("", s)
    s = html.unescape(s).replace("\xa0", " ")
    s = clean(s)
    return re.sub(r"\n{3,}", "\n\n", s)


def clean_title(title, poet):
    title = clean(html.unescape(str(title or "")).replace("\xa0", " "))
    if poet and title.endswith("/ " + poet):
        title = title[: -len(poet) - 2].rstrip(" /")
    elif " / " in title:  # poet spelled differently in the title
        title = title.rsplit(" / ", 1)[0]
    return " ".join(title.split())


def read_rows(path):
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().upper() for h in next(rows)]
    try:
        cols = [header.index(c) for c in ("POET", "TITLE", "TEXT")]
    except ValueError:
        sys.exit(f"Expected columns POET, TITLE, TEXT; found {header}")
    for row in rows:
        poet, title, text = (row[i] if i < len(row) else None for i in cols)
        poet = " ".join(clean(str(poet or "")).split())
        yield poet, clean_title(title, poet), html_to_text(text)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("xlsx")
    ap.add_argument("--dry-run", action="store_true", help="parse and count, don't write")
    ap.add_argument("--show", type=int, default=0, metavar="N", help="print the first N parsed poems")
    args = ap.parse_args()

    shayri.init_db()
    db = shayri.connect()
    existing = {fold(" ".join(r[0].split())) for r in db.execute("SELECT body FROM poems")}
    added = dup = empty = 0
    with db:
        for n, (poet, title, body) in enumerate(read_rows(args.xlsx), 1):
            if n <= args.show:
                print(f"--- {title} / {poet}\n{body}\n")
            if not body:
                empty += 1
                continue
            key = fold(" ".join(body.split()))
            if key in existing:
                dup += 1
                continue
            existing.add(key)
            if not args.dry_run:
                shayri.save_poem(db, {"title": title, "poet": poet, "body": body,
                                      "tags": "", "notes": "", "source": "", "favorite": 0})
            added += 1
            if n % 2000 == 0:
                print(f"  {n} rows read...", file=sys.stderr)
        if args.dry_run:
            db.rollback()
    verb = "Would import" if args.dry_run else "Imported"
    print(f"{verb} {added} poem(s); skipped {dup} duplicate(s), {empty} empty row(s). DB: {shayri.DB_PATH}")


if __name__ == "__main__":
    main()
